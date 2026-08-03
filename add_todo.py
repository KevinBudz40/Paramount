"""
add_todo.py — Append tasks to Kevin_Todo.xlsx without rebuilding it.

Usage (interactive):
    python add_todo.py

Usage (one-liner, all args on CLI):
    python add_todo.py --task "Do something" --workstream "AR / Integrations" \
        --source "Mtg 8/4" --due "2026-08-07" --notes "Details here"

Optional flags:
    --status    Open | In Progress | Done | Blocked   (default: Open)
    --xlsx      Path to the xlsx file                 (default: C:\\Paramount\\local_docs\\Kevin_Todo.xlsx)
    --no-prompt Skip all prompts; require --task at minimum
"""

import argparse
import datetime
import sys
import os

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("ERROR: openpyxl not installed. Run: pip install openpyxl")

XLSX_DEFAULT = r"C:\Paramount\local_docs\Kevin_Todo.xlsx"

# ── style constants (must match build_todo_v2.py) ────────────────────────────
ROW_ALT  = PatternFill("solid", fgColor="EEF2F7")
ROW_EVEN = PatternFill("solid", fgColor="FFFFFF")
STATUS_COLORS = {
    "Open":        PatternFill("solid", fgColor="FFD966"),
    "In Progress": PatternFill("solid", fgColor="9DC3E6"),
    "Done":        PatternFill("solid", fgColor="A9D18E"),
    "Blocked":     PatternFill("solid", fgColor="FF7373"),
}
THIN   = Side(border_style="thin", color="C9C9C9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
VALID_STATUSES = list(STATUS_COLORS.keys())


def _prompt(label, default=None, choices=None):
    """Prompt user with an optional default. Return stripped input."""
    hint = f" [{default}]" if default else ""
    if choices:
        hint = f" ({'/'.join(choices)}){hint}"
    raw = input(f"  {label}{hint}: ").strip()
    if not raw and default is not None:
        return default
    return raw


def append_task(xlsx_path, status, task, workstream, source, date_added, due_date, notes):
    """Load existing workbook, append one formatted row, save."""
    if not os.path.exists(xlsx_path):
        sys.exit(f"ERROR: file not found: {xlsx_path}")

    wb = load_workbook(xlsx_path)
    ws = wb.active

    # find the next empty row
    next_row = ws.max_row + 1

    fill = ROW_ALT if next_row % 2 == 0 else ROW_EVEN
    values = [status, task, workstream, source, date_added, due_date, notes]

    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(col_idx in (2, 7)))
        if col_idx == 1:
            cell.fill = STATUS_COLORS.get(status, fill)
            cell.alignment = Alignment(horizontal="center", vertical="top")
            cell.font = Font(bold=True, name="Calibri", size=10)
        else:
            cell.fill = fill
            cell.font = Font(name="Calibri", size=10)

    ws.row_dimensions[next_row].height = 52

    # extend auto-filter to cover new row
    num_cols = len(values)
    ws.auto_filter.ref = f"A1:{get_column_letter(num_cols)}{next_row}"

    wb.save(xlsx_path)
    return next_row - 1  # task number (data rows start at 2)


def main():
    parser = argparse.ArgumentParser(description="Append a task to Kevin_Todo.xlsx")
    parser.add_argument("--task",        help="Task description")
    parser.add_argument("--workstream",  help="Workstream label")
    parser.add_argument("--source",      help="Source (meeting/email/etc.)")
    parser.add_argument("--due",         help="Due date (YYYY-MM-DD)")
    parser.add_argument("--notes",       help="Notes / context")
    parser.add_argument("--status",      default="Open",
                        help=f"Status ({'/'.join(VALID_STATUSES)}), default: Open")
    parser.add_argument("--date-added",  help="Date added (YYYY-MM-DD), default: today")
    parser.add_argument("--xlsx",        default=XLSX_DEFAULT, help="Path to xlsx file")
    parser.add_argument("--no-prompt",   action="store_true",
                        help="Skip interactive prompts; fail if required args missing")
    args = parser.parse_args()

    today = datetime.date.today().isoformat()

    if args.no_prompt:
        # non-interactive mode — require at minimum --task
        if not args.task:
            sys.exit("ERROR: --task is required in --no-prompt mode")
        task        = args.task
        workstream  = args.workstream or ""
        source      = args.source     or ""
        due_date    = args.due        or ""
        notes       = args.notes      or ""
        status      = args.status
        date_added  = args.date_added or today
    else:
        # interactive mode — pre-fill from CLI args, prompt for anything missing
        print("\n── Add a new todo ──────────────────────────────────────────")
        task       = args.task       or _prompt("Task (required)")
        if not task:
            sys.exit("ERROR: task cannot be empty")
        workstream = args.workstream or _prompt("Workstream", default="AR / Integrations")
        source     = args.source     or _prompt("Source (meeting/email/etc.)")
        due_date   = args.due        or _prompt("Due date (YYYY-MM-DD)", default="")
        notes      = args.notes      or _prompt("Notes", default="")
        status     = args.status
        if status not in VALID_STATUSES:
            status = _prompt("Status", default="Open", choices=VALID_STATUSES)
        date_added = args.date_added or today

    if status not in VALID_STATUSES:
        sys.exit(f"ERROR: invalid status '{status}'. Choose from: {', '.join(VALID_STATUSES)}")

    task_num = append_task(
        xlsx_path  = args.xlsx,
        status     = status,
        task       = task,
        workstream = workstream,
        source     = source,
        date_added = date_added,
        due_date   = due_date,
        notes      = notes,
    )
    print(f"\nOK Task #{task_num} added to {args.xlsx}")
    ellipsis = "..." if len(task) > 70 else ""
    print(f"  [{status}] {task[:70]}{ellipsis}")


if __name__ == "__main__":
    main()
