"""
scan_todos.py  — Compare a list of extracted action items against the existing
                 Kevin_Todo.xlsx and append net-new ones.

This script is the EXCEL LAYER — it doesn't scan sources itself.
Sources (Tactiq, email, Slack) are read by Claude Code in-session, which
extracts action items and pipes them here as JSON.

Usage (called by Claude Code):
    python scan_todos.py --input tasks.json          # append from file
    python scan_todos.py --input tasks.json --dry-run  # preview only
    python scan_todos.py --print-existing            # dump current task list

Input JSON format (array of objects):
    [
      {
        "task":       "Do the thing",
        "workstream": "AR / Integrations",
        "source":     "Meeting: Avalara UIP 2026-08-04",
        "due_date":   "2026-08-07",
        "notes":      "Context here",
        "status":     "Open"           (optional, default Open)
      },
      ...
    ]

Cursor file: C:\\GIT\\todo_cursor.json  (read/written by Claude Code scan workflow)
Excel:       C:\\Paramount\\local_docs\\Kevin_Todo.xlsx
"""

import argparse
import datetime
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("ERROR: pip install openpyxl")

XLSX_PATH    = Path(r"C:\Paramount\local_docs\Kevin_Todo.xlsx")
CURSOR_FILE  = Path(__file__).parent / "todo_cursor.json"

STATUS_COLORS = {
    "Open":        PatternFill("solid", fgColor="FFD966"),
    "In Progress": PatternFill("solid", fgColor="9DC3E6"),
    "Done":        PatternFill("solid", fgColor="A9D18E"),
    "Blocked":     PatternFill("solid", fgColor="FF7373"),
}
ROW_ALT  = PatternFill("solid", fgColor="EEF2F7")
ROW_EVEN = PatternFill("solid", fgColor="FFFFFF")
THIN     = Side(border_style="thin", color="C9C9C9")
BORDER   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ── normalise for dedup ───────────────────────────────────────────────────────
def _norm(text):
    return re.sub(r"[^a-z0-9 ]", "", text.lower())


def _is_duplicate(new_task, existing_norms, threshold=0.55):
    """True if new_task overlaps heavily with any existing task."""
    nw = set(_norm(new_task).split())
    if len(nw) < 3:
        return _norm(new_task) in existing_norms
    for ex in existing_norms:
        ew = set(ex.split())
        overlap = len(nw & ew) / max(len(nw), 1)
        if overlap >= threshold:
            return True
    return False


def load_existing_norms(xlsx_path):
    """Return normalised task strings already in the Excel."""
    if not xlsx_path.exists():
        return set()
    wb = load_workbook(xlsx_path)
    ws = wb.active
    norms = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1]:
            norms.add(_norm(str(row[1])))
    return norms


def print_existing(xlsx_path):
    """Dump current tasks for reference."""
    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}")
        return
    wb = load_workbook(xlsx_path)
    ws = wb.active
    print(f"\nExisting tasks in {xlsx_path.name}:")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1]:
            print(f"  [{row[0]}] {str(row[1])[:80]}")


# ── Excel append ──────────────────────────────────────────────────────────────
def append_tasks(xlsx_path, tasks):
    wb = load_workbook(xlsx_path)
    ws = wb.active
    today = datetime.date.today().isoformat()

    for task in tasks:
        row_idx = ws.max_row + 1
        fill    = ROW_ALT if row_idx % 2 == 0 else ROW_EVEN
        status  = task.get("status", "Open")
        values  = [
            status,
            task.get("task", ""),
            task.get("workstream", ""),
            task.get("source", ""),
            task.get("date_added", today),
            task.get("due_date", ""),
            task.get("notes", ""),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(col_idx in (2, 7)))
            if col_idx == 1:
                cell.fill      = STATUS_COLORS.get(status, fill)
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font      = Font(bold=True, name="Calibri", size=10)
            else:
                cell.fill = fill
                cell.font = Font(name="Calibri", size=10)
        ws.row_dimensions[row_idx].height = 52

    ws.auto_filter.ref = f"A1:{get_column_letter(7)}{ws.max_row}"
    wb.save(xlsx_path)


# ── cursor helpers ────────────────────────────────────────────────────────────
def load_cursor():
    if CURSOR_FILE.exists():
        return json.loads(CURSOR_FILE.read_text())
    return {}


def save_cursor(cursor):
    CURSOR_FILE.write_text(json.dumps(cursor, indent=2))


# ── main ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="Append Claude-extracted action items to Kevin_Todo.xlsx"
    )
    parser.add_argument("--input",          help="JSON file of extracted tasks")
    parser.add_argument("--dry-run",        action="store_true",
                        help="Print what would be added without writing")
    parser.add_argument("--print-existing", action="store_true",
                        help="Dump existing task list and exit")
    parser.add_argument("--update-cursor",  help="ISO timestamp to write into cursor (source:ts)")
    args = parser.parse_args()

    if args.print_existing:
        print_existing(XLSX_PATH)
        return

    if args.update_cursor:
        # format: "tactiq:2026-08-04T12:00:00+00:00"
        source, ts = args.update_cursor.split(":", 1)
        cursor = load_cursor()
        cursor[source] = ts
        save_cursor(cursor)
        print(f"Cursor updated: {source} -> {ts}")
        return

    if not args.input:
        parser.print_help()
        return

    input_path = Path(args.input)
    if not input_path.exists():
        sys.exit(f"ERROR: {input_path} not found")

    all_tasks = json.loads(input_path.read_text(encoding="utf-8"))
    if not isinstance(all_tasks, list):
        sys.exit("ERROR: input JSON must be a list of task objects")

    existing = load_existing_norms(XLSX_PATH)
    net_new  = []

    for t in all_tasks:
        task_text = t.get("task", "").strip()
        if not task_text:
            continue
        if _is_duplicate(task_text, existing):
            print(f"  [skip dup] {task_text[:70]}")
            continue
        net_new.append(t)
        existing.add(_norm(task_text))

    print(f"\n{len(net_new)} net-new tasks (of {len(all_tasks)} candidates)")
    for i, t in enumerate(net_new, 1):
        print(f"  {i}. [{t.get('workstream','')}] {t['task'][:70]}")
        if t.get("due_date"):
            print(f"       Due: {t['due_date']}")
        if t.get("notes"):
            print(f"       Notes: {t['notes'][:80]}")

    if net_new and not args.dry_run:
        if not XLSX_PATH.exists():
            sys.exit(f"ERROR: {XLSX_PATH} not found")
        append_tasks(XLSX_PATH, net_new)
        print(f"\nAppended {len(net_new)} tasks to {XLSX_PATH}")
    elif args.dry_run:
        print("\n[dry-run] Nothing written.")
    else:
        print("\nNothing new to add.")


if __name__ == "__main__":
    main()
